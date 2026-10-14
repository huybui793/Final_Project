import os
import json
from io import BytesIO
from google import genai
from dotenv import load_dotenv
from pptx import Presentation

load_dotenv()
client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))


def generate_slides(context, topic, num_slides=6):
    prompt = f"""Dựa vào tài liệu sau, hãy soạn bộ slide thuyết trình về chủ đề "{topic}" gồm {num_slides} slide.
Slide đầu là slide giới thiệu, slide cuối là slide tóm tắt. Mỗi slide có tiêu đề ngắn và 3 đến 4 ý, chi tiết không quá 100 từ.
Chỉ dùng thông tin có trong tài liệu.

Tài liệu:
{context}

Trả về CHÍNH XÁC theo định dạng JSON, không thêm text nào khác:
[
  {{"title": "...", "bullets": ["...", "...", "..."]}}
]
"""
    response = client.models.generate_content(model="gemini-2.5-flash", contents=prompt)
    text = response.text.strip()
    if text.startswith("```"):
        text = text.split("```")[1]
        if text.startswith("json"):
            text = text[4:]
    return json.loads(text.strip())


def build_pptx(slides, deck_title):
    prs = Presentation()

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = deck_title
    cover.placeholders[1].text = "Trợ lý học tập Lịch sử AI"

    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s["title"]
        body = slide.placeholders[1].text_frame
        body.text = s["bullets"][0] if s["bullets"] else ""
        for bullet in s["bullets"][1:]:
            body.add_paragraph().text = bullet

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()